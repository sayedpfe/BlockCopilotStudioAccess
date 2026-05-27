# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu, Inches
prs = Presentation("presentation-cps-lockdown-2026-05-27.pptx")
SW, SH = prs.slide_width, prs.slide_height
M = Inches(0.4)
viol = 0
for si, slide in enumerate(prs.slides,1):
    for shp in slide.shapes:
        try:
            l,t,w,h = shp.left, shp.top, shp.width, shp.height
        except Exception:
            continue
        if l is None or t is None or w is None or h is None: continue
        r, b = l+w, t+h
        msgs=[]
        if l < 0: msgs.append(f"left<0 ({Emu(l).inches:.2f}in)")
        if t < 0: msgs.append(f"top<0 ({Emu(t).inches:.2f}in)")
        if r > SW: msgs.append(f"right>{Emu(SW).inches:.2f} ({Emu(r).inches:.2f}in)")
        if b > SH: msgs.append(f"bottom>{Emu(SH).inches:.2f} ({Emu(b).inches:.2f}in)")
        if msgs:
            viol+=1
            nm = shp.shape_type
            txt = (shp.text_frame.text[:30] if shp.has_text_frame else "")
            print(f"  S{si}: {nm} '{txt}' -> {', '.join(msgs)}")
print(f"BOUNDS VIOLATIONS: {viol}")
